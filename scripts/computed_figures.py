#!/usr/bin/env python3
"""Which site-capacity figures were computed rather than stated, and how.

ROADMAP #248: a figure we assemble is not a figure a source states.
"Wind Generation of 3 no. 900kW turbines" is stored as 2.7 MW, and the
2.7 appears nowhere in the quote beneath it. The arithmetic is right;
the page renders it exactly like a disclosed figure, and a reporter
checking it finds the components and no total. Before a derivation
record can be designed, the question is what the derivations ARE: how
many are a unit count times a rating, how many a sum, and how many
something no arithmetic on the quote reaches.

This is the measurement, kept as a script because the number moves
with the corpus and a number copied out of a session is a number
nobody can re-run (README, on the null-capacity sweep). It reads every
`site_capacity` figure on a live site whose figures stand (migration
034), asks whether any number in its quote states the value — in MW, in
kW, or as the original value — and classes the rest:

  A. a product of operands in the quote: a unit count times a rating,
     every fleet the quote names summed, or two stated numbers
  B. a sum of two or three stated figures
  C. a stated figure scaled (80%, ×5 …) or the midpoint of a range —
     an inference dressed as a figure, not a derivation
  D. nothing arithmetic on the quote reaches it: an operand taken from
     the passage beyond the quote, a count in words, a substrate too
     broken to read ("BSOOMW", "5|0MW"), or a number invented

The quote is repaired before matching — decimal commas ("1720,71 kW"),
digit spacing ("1 250", "3 .3"), and OCR'd letters in numbers next to
a unit ("SMW", "7OMW", "ISMW") — because a figure the substrate states
badly is stated, and counting it as computed would inflate the class
the record is for. What survives repair and still cannot be reached is
the honest residue, and it is listed in full.

Measured 2026-09-10: 10,500 figures, 206 computed (2.0%); A 75, B 7,
C 39, D 85. Re-run rather than quote.

    scripts/computed_figures.py            # writes data/reports/computed_figures_<date>.md
    scripts/computed_figures.py --print    # to stdout as well
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from dotenv import load_dotenv  # noqa: E402

load_dotenv(ROOT / ".env")

from dcp import db  # noqa: E402

SQL = """
WITH latest AS (
  SELECT DISTINCT ON (finding_id) finding_id, application_id, verdict,
         quantity_type, value_mw, value_original, unit_original, model, id
  FROM power_adjudication
  ORDER BY finding_id, (verdict = 'unclear'), inserted_at DESC, id DESC)
SELECT la.finding_id, la.quantity_type, la.value_mw, la.value_original,
       f.evidence_text, s.site_key, a.application_ref, la.model,
       f.document_id, f.evidence_page, la.id
FROM latest la
JOIN findings f ON f.id = la.finding_id
JOIN applications a ON a.id = la.application_id
JOIN site_members sm ON sm.application_id = la.application_id
     AND sm.retired_at IS NULL
     AND sm.figure_standing <> 'not_dc_excluded'
JOIN sites s ON s.id = sm.site_id AND s.retired_at IS NULL
WHERE la.verdict = 'site_capacity' AND la.value_mw IS NOT NULL
  AND f.evidence_text IS NOT NULL
ORDER BY s.site_key, la.finding_id
"""

from dcp.derivation import DERIVATION_VERSION, classify, derive, stated  # noqa: E402


REVIEW_DIR = ROOT / "data" / "computed_figures_review"
# The four families a site's figure box shows; a cooling or storage
# rating never heads a page, so it is not "the figure a site shows".
BOX_FAMILIES = ("it_load", "total_site", "grid_connection", "onsite_generation")
REVIEW_HEADERS = ["class", "site_key", "site_name", "application_ref",
                  "quantity_type", "value_mw", "value_original",
                  "the site's shown figure?", "quote", "document (our copy)",
                  "source url", "page", "reader", "finding_id",
                  "decision (keep / unclear / correct to …)", "notes"]


def write_derivations(rows) -> int:
    """The A and B figures' derivations, recorded beside their adjudications
    (migration 035). Idempotent: a row already carrying one under this
    version is skipped, so the pass can run after every adjudication
    batch — though from 2026-09-15 the adjudication scripts record the
    derivation at write, and this is the backfill for what preceded it."""
    from dcp import derivation
    written = seen = 0
    with db.connect() as conn, conn.cursor() as cur:
        for (fid, qt, mw, vo, quote, key, ref, model, doc_id, page, adj_id) in rows:
            mw = float(mw)
            if stated(mw, vo, quote):
                continue
            d = derive(mw, quote)
            if d is None:
                continue
            seen += 1
            if derivation.record(cur, adjudication_id=adj_id, finding_id=fid,
                                 value_mw=mw, d=d):
                written += 1
        conn.commit()
    print(f"derivations: {seen} A and B figures, {written} recorded under "
          f"{DERIVATION_VERSION} ({seen - written} already carried one)")
    return written


def write_review(computed, rows) -> Path:
    """The C and D figures as a workbook for a person to settle by hand,
    one row per figure, the same shape as the operator-pages review:
    every column the decision needs beside two empty ones for the
    decision and the reason. Luke chose this over `unclear` on
    2026-09-10 — the reader loses little by withdrawing a mangled
    cooling rating, but the workbook keeps the lead. Rewritten on every
    run from the classification, so a decision column is folded back
    into the record (a person's adjudication row) before the next run,
    never edited here."""
    import openpyxl
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from dcp import drive
    picked = [c for c in computed if c[0][:1] in ("C", "D")]
    by_fid = {r[0]: r for r in rows}
    # The figure a site's box shows is the largest standing figure of
    # its family, which is what the reader's rollups take.
    with db.connect() as conn, conn.cursor() as cur:
        # The generation rung shows standby-shaped plant only (the same
        # filter as site_cohorts.SITE_FIGURES_SQL): plant adjudicated
        # prime_combustion, renewable or storage, or not generation at
        # all, never heads a page.
        cur.execute("""
            SELECT finding_id FROM (
              SELECT DISTINCT ON (finding_id) finding_id, figure_basis, plant_type
              FROM generation_adjudication
              ORDER BY finding_id, inserted_at DESC, id DESC) g
            WHERE coalesce(g.figure_basis, '') = 'not_generation'
               OR coalesce(g.plant_type, '') IN ('prime_combustion', 'renewable', 'storage')""")
        not_shown_generation = {r[0] for r in cur.fetchall()}
        shown = {}
        for (fid, qt, mw, *_rest) in rows:
            if qt not in BOX_FAMILIES:
                continue
            if qt == "onsite_generation" and fid in not_shown_generation:
                continue
            key = by_fid[fid][5], qt
            shown[key] = max(shown.get(key, 0.0), float(mw))
        cur.execute("SELECT site_key, display_name FROM sites WHERE retired_at IS NULL")
        names = dict(cur.fetchall())
        doc_ids = sorted({c[7] for c in picked if c[7]})
        cur.execute("""
            SELECT DISTINCT ON (d.id) d.id, d.url, f.file_id
            FROM documents d
            LEFT JOIN document_drive_files f ON f.document_id = d.id
            WHERE d.id = ANY(%s)
            ORDER BY d.id, f.recorded_at DESC NULLS LAST""", (doc_ids,))
        docs = {i: (u, fid) for i, u, fid in cur.fetchall()}
    out_rows = []
    for (k, key, ref, qt, mw, model, fid, doc_id, page, q) in picked:
        vo = by_fid[fid][3]
        unit = ""
        url, file_id = docs.get(doc_id, (None, None))
        out_rows.append([
            k[:1], key, names.get(key, ""), ref, qt, mw,
            (f"{vo:g}" if vo is not None else ""),
            "yes" if abs(shown.get((key, qt), 0.0) - mw) < 1e-9 else "",
            q, drive.file_url(file_id) if file_id else "", url or "", page or "",
            model, fid, "", ""])
    out_rows.sort(key=lambda r: (r[7] != "yes", r[0], r[1], -float(r[5])))
    REVIEW_DIR.mkdir(parents=True, exist_ok=True)
    with (REVIEW_DIR / "computed_figures_review.csv").open("w", newline="",
                                                           encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(REVIEW_HEADERS)
        w.writerows(out_rows)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "computed_figures_review"
    ws.append(REVIEW_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in out_rows:
        ws.append(r)
    widths = [6, 26, 24, 24, 16, 9, 10, 9, 60, 34, 34, 6, 18, 10, 22, 30]
    for i, wd in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    dest = REVIEW_DIR / "computed_figures_review.xlsx"
    wb.save(dest)
    print(f"wrote {dest}: {len(out_rows)} C and D figures for review, "
          f"{sum(1 for r in out_rows if r[7] == 'yes')} of them the figure a site shows")
    return dest


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--print", action="store_true", help="also print the report")
    ap.add_argument("--out", type=Path, default=None)
    ap.add_argument("--derive", action="store_true",
                    help="record a derivation (migration 035) beside every "
                         "A or B adjudication that lacks one under "
                         f"{DERIVATION_VERSION}; append-only, a no-op on re-run")
    ap.add_argument("--review", action="store_true",
                    help="also write the C and D figures as a review workbook "
                         "under data/computed_figures_review/ for a person to "
                         "settle by hand (Luke, 2026-09-10)")
    args = ap.parse_args()

    with db.connect() as conn, conn.cursor() as cur:
        cur.execute(SQL)
        rows = cur.fetchall()
    computed = []
    for (fid, qt, mw, vo, quote, key, ref, model, doc_id, page, adj_id) in rows:
        mw = float(mw)
        if stated(mw, vo, quote):
            continue
        computed.append((classify(mw, quote), key, ref, qt, mw, model, fid, doc_id,
                         page, " ".join((quote or "").split())))
    computed.sort()
    kinds = Counter(c[0] for c in computed)
    stamp = dt.datetime.now(dt.timezone.utc)
    out = [f"# Computed site-capacity figures — {stamp:%Y-%m-%d %H:%M} UTC", "",
           f"{len(rows):,} site-capacity figures on live sites whose figures stand; "
           f"**{len(computed):,} ({100 * len(computed) / max(len(rows), 1):.1f}%) hold "
           f"a value no number in their quote states**, in megawatts, kilowatts or "
           f"as the original value, read as written and as repaired for decimal "
           f"commas, digit spacing and OCR'd digits.", "",
           f"Across {len({c[1] for c in computed})} sites; by reader: "
           + ", ".join(f"{m} {n}" for m, n in Counter(c[5] for c in computed).most_common()),
           "", "| class | figures |", "|---|---:|"]
    out += [f"| {k} | {n} |" for k, n in sorted(kinds.items())]
    out += ["", "Class A and B are derivations a record can carry: the operands are "
                "in the quote and the operation is stated. Class C is an inference — "
                "a scaling or a midpoint is a choice, not arithmetic the source made. "
                "Class D is the residue to read: an operand taken from the passage "
                "beyond the quote, a count the fleet pattern does not parse, a "
                "substrate too broken to repair, or a number with no source.", ""]
    for k in sorted(kinds):
        out += [f"## {k} ({kinds[k]})", ""]
        for (_k, key, ref, qt, mw, model, fid, doc_id, page, q) in computed:
            if _k != k:
                continue
            out.append(f"- **{mw:g} MW** {qt} · {key} · {ref} · finding {fid}"
                       + (f" · document {doc_id}" if doc_id else "")
                       + (f", page {page}" if page else "")
                       + f" · {model}\n  > {q[:300]}")
        out.append("")
    report = "\n".join(out)
    dest = args.out or (ROOT / "data" / "reports" / f"computed_figures_{stamp:%Y-%m-%d}.md")
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text(report, encoding="utf-8")
    if args.review:
        write_review(computed, rows)
    if args.derive:
        write_derivations(rows)
    print(f"wrote {dest}: {len(computed)} of {len(rows)} figures computed; "
          + ", ".join(f"{k[:1]} {n}" for k, n in sorted(kinds.items())))
    if args.print:
        print(report)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
